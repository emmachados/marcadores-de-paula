import csv
import json
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from transcript_desk.annotations import create_annotation, create_range_annotation, create_time_annotation, is_current
from transcript_desk.exports import export_result
from transcript_desk import turn_edits


def test_untimed_speaker_export_and_filtered_concordance(tmp_path, result):
    result['segments'].append({'start': 5, 'end': 6, 'speaker': 'SPEAKER_00', 'text': 'Last question.'})
    result['annotations'] = [create_annotation(result, 0, 0, 3, 'Other'),
                             create_annotation(result, 1, 0, 3, 'Other', status='suggested')]
    export_result(result, tmp_path, ['TXT', 'DOCX', 'JSON'], options={
        'timestamps': False, 'grouped': True, 'paragraphs': False, 'speakers': ['SPEAKER_00'],
        'statuses': ['confirmed'], 'analysis_tables': True, 'marked_passages': True})
    text = (tmp_path / 'transcript.txt').read_text()
    assert '00:00' not in text and 'Last question.' in text
    assert 'Participant' not in text
    assert len(json.loads((tmp_path / 'transcript.json').read_text())['segments']) == 2
    assert len(list(csv.DictReader((tmp_path / 'concordance.csv').open(encoding='utf-8-sig')))) == 1
    assert any('Last question.' in p.text for p in Document(tmp_path / 'transcript.docx').paragraphs)


@pytest.mark.parametrize('timestamps, paragraphs', [(False, False), (False, True), (True, True)])
def test_chronological_docx_follows_text_layout(tmp_path, result, timestamps, paragraphs):
    result['segments'].insert(2, {'start': 3.5, 'end': 3.9, 'speaker': 'SPEAKER_01', 'text': 'Second answer.'})
    export_result(result, tmp_path, ['TXT', 'DOCX'], options={
        'timestamps': timestamps, 'grouped': False, 'paragraphs': paragraphs})
    body = Document(tmp_path / 'transcript.docx').paragraphs[2:]
    assert [p.text for p in body if p.style.name.startswith('Heading')] == ['Interviewer', 'Participant', 'Unassigned']
    parts = []
    for p in body:
        if p.style.name.startswith('Heading'):
            parts.append(('\n\n' if parts else '') + p.text + '\n')
        else:
            parts.append(('' if parts[-1].endswith('\n') else '\n\n') + p.text)
    assert ''.join(parts) + '\n' == (tmp_path / 'transcript.txt').read_text()


def test_annotation_tiers_handle_overlap_without_expanding_spans(tmp_path, result):
    first = create_annotation(result, 0, 0, 3, 'Other')
    second = create_annotation(result, 0, 4, 8, 'Other')
    result['annotations'] = [first, second]
    export_result(result, tmp_path, ['EAF', 'TextGrid'], options={'annotation_tiers': True})
    tree = ET.parse(tmp_path / 'transcript.eaf')
    assert len(tree.findall('.//ALIGNABLE_ANNOTATION')) == 5
    names = [tier.attrib['TIER_ID'] for tier in tree.findall('.//TIER')]
    assert len([name for name in names if name.startswith('annotation_')]) == 2
    assert 'Other [2]' in (tmp_path / 'transcript.TextGrid').read_text()


def test_multi_turn_and_time_annotations_survive_split_merge(result):
    result['segments'][0]['text'] = 'First and second'
    result['annotations'] = [create_range_annotation(result, 0, 1, 'Other'),
                             create_time_annotation(result, 0, 1.3, 1.5, 'Breath')]
    turn_edits.split_turn(result, 0, 6, 1.1)
    multi, event = result['annotations']
    assert len(multi['spans']) == 3 and is_current(multi, result)
    assert event['segment'] == 1 and event['start'] == 1.3
    turn_edits.merge_turn(result, 0)
    assert is_current(multi, result)
    assert event['segment'] == 0 and event['start'] == 1.3


def test_filtered_annotation_rows_keep_correct_turn_ids(tmp_path, result):
    result['annotations'] = [create_annotation(result, 1, 0, 3, 'Other')]
    export_result(result, tmp_path, ['JSON'], options={'speakers': ['SPEAKER_01']})
    exported = json.loads((tmp_path / 'transcript.json').read_text())
    assert exported['annotations'][0]['segment'] == 0
    assert exported['annotations'][0]['turn_id'] == exported['segments'][0]['turn_id']
    assert is_current(exported['annotations'][0], exported)
