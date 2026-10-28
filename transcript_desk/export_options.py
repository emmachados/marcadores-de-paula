"""Export selections, annotation tiers and counting tables."""
import copy
import csv
import json
from collections import defaultdict

from transcript_desk.codebook import annotation_name
from transcript_desk.exports import speaker_name, timestamp


def selected_result(result, options):
    revised = copy.deepcopy(result)
    speakers = options.get('speakers')
    rows = [i for i, s in enumerate(revised['segments']) if speakers is None or s.get('speaker') in speakers]
    row_map = {old: new for new, old in enumerate(rows)}
    kept_ids = {revised['segments'][row].get('turn_id') for row in rows}
    annotations = []
    statuses, labels = options.get('statuses'), options.get('labels')
    for annotation in revised.get('annotations', []):
        if annotation['segment'] not in row_map:
            continue
        if statuses is not None and annotation['status'] not in statuses:
            continue
        if labels is not None and annotation.get('label_id', annotation['label']) not in labels:
            continue
        if any(span['turn_id'] not in kept_ids for span in annotation.get('spans', [])):
            continue
        annotation['segment'] = row_map[annotation['segment']]
        annotations.append(annotation)
    revised['segments'] = [revised['segments'][row] for row in rows]
    revised['annotations'] = annotations
    revised.setdefault('metadata', {})['export_selection'] = {
        'source_rows': rows, 'statuses': statuses, 'labels': labels,
        'text_layout': 'speaker' if options.get('grouped') else 'chronological',
        'timestamps_in_text': options.get('timestamps', True)}
    return revised


def annotation_tier_segments(result):
    """Put overlapping annotations on separate tiers; each keeps its own span."""
    by_turn = {s.get('turn_id'): s for s in result['segments']}
    intervals = defaultdict(list)
    for annotation in result.get('annotations', []):
        if annotation['status'] == 'needs review':
            continue
        label = annotation_name(result, annotation)
        spans = annotation.get('spans', [])
        if spans:
            bounds = [(by_turn[span['turn_id']], by_turn[span['turn_id']]['start'], by_turn[span['turn_id']]['end'], span['quote']) for span in spans]
        else:
            bounds = [(result['segments'][annotation['segment']], annotation['start'], annotation['end'], annotation['quote'])]
        for source, start, end, quote in bounds:
            name = f"{speaker_name(source.get('speaker'), result.get('speaker_names', {}))} · {label}"
            text = quote or label
            if annotation.get('note'):
                text += ' [' + annotation['note'] + ']'
            intervals[name].append((start, end, text))
    rows, names = [], {}
    for group, entries in sorted(intervals.items()):
        lanes = []
        for start, end, text in sorted(entries):
            lane = next((i for i, last_end in enumerate(lanes) if last_end <= start), len(lanes))
            if lane == len(lanes):
                lanes.append(end)
            else:
                lanes[lane] = end
            identity = f"annotation_{len(names)}"
            lane_name = group + (f' [{lane + 1}]' if lane else '')
            identity = next((key for key, value in names.items() if value == lane_name), identity)
            names[identity] = lane_name
            rows.append({'start': start, 'end': end, 'text': text, 'speaker': identity})
    return rows, names


def add_document_content(document, result, options):
    if options.get('grouped'):
        from transcript_desk.text_tools import reading_view, speaker_ids
        for identity in speaker_ids(result):
            document.add_heading(speaker_name(identity, result.get('speaker_names', {})), level=1)
            text, _ = reading_view(result, speaker=identity, all_speakers=False,
                                  timestamps=options.get('timestamps', True), paragraphs=options.get('paragraphs', True))
            for paragraph in text.split('\n', 1)[-1].split('\n\n'):
                document.add_paragraph(paragraph)
    else:
        from docx.enum.text import WD_COLOR_INDEX
        # Same layout as the TXT and the preview: a heading only where the speaker changes.
        rows = sorted(enumerate(result['segments']), key=lambda item: (item[1]['start'], item[1]['end'], item[0]))
        previous, paragraph = object(), None
        for row, segment in rows:
            if segment.get('speaker') != previous:
                document.add_heading(speaker_name(segment.get('speaker'), result.get('speaker_names', {})), level=2)
                paragraph = document.add_paragraph()
            elif options.get('paragraphs', True):
                paragraph = document.add_paragraph()
            else:
                paragraph.add_run(' ')
            previous = segment.get('speaker')
            if options.get('timestamps', True):
                paragraph.add_run(f"[{timestamp(segment['start'], '.')} – {timestamp(segment['end'], '.')}] ")
            annotations = [a for a in result.get('annotations', []) if a['status'] != 'needs review' and
                           (a['segment'] == row or any(span['turn_id'] == segment.get('turn_id') for span in a.get('spans', [])))]
            if options.get('marked_passages') and annotations:
                boundaries = {0, len(segment['text'])}
                ranges = []
                for annotation in annotations:
                    anchors = ([span for span in annotation['spans'] if span['turn_id'] == segment.get('turn_id')]
                               if annotation.get('spans') else [annotation])
                    for anchor in anchors:
                        if annotation.get('anchor_type') == 'time':
                            continue
                        start, end = anchor['start_char'], anchor['end_char']
                        if 0 <= start < end <= len(segment['text']):
                            boundaries.update((start, end))
                            ranges.append((start, end))
                boundaries = sorted(boundaries)
                for start, end in zip(boundaries, boundaries[1:]):
                    run = paragraph.add_run(segment['text'][start:end])
                    if any(left < end and right > start for left, right in ranges):
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            else:
                paragraph.add_run(segment['text'])
    if options.get('marked_passages') and result.get('annotations'):
        document.add_heading('Annotations', level=1)
        for annotation in result['annotations']:
            document.add_paragraph(f"{annotation_name(result, annotation)} · {annotation['status']} · {annotation['quote']}")
            if annotation.get('note'):
                document.add_paragraph(annotation['note'])


def _write_csv(path, fields, rows):
    def safe(value):
        return "'" + value if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')) else value
    with path.open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.writer(stream)
        writer.writerow(fields)
        for row in rows:
            writer.writerow([safe(value) for value in row])


def export_analysis(result, folder):
    from transcript_desk.text_tools import speaker_statistics
    concordance, counts = [], defaultdict(int)
    by_turn = {s.get('turn_id'): s for s in result['segments']}
    for annotation in result.get('annotations', []):
        label = annotation_name(result, annotation)
        current = annotation['status'] != 'needs review' and annotation.get('anchor_type') != 'time'
        anchors = annotation.get('spans') or [annotation]
        counted = set()
        for index, anchor in enumerate(anchors):
            segment = by_turn.get(anchor.get('turn_id'), result['segments'][annotation['segment']])
            name = speaker_name(segment.get('speaker'), result.get('speaker_names', {}))
            identity = segment.get('speaker') or ''
            if identity not in counted:
                counts[(identity, name, label, annotation['status'])] += 1
                counted.add(identity)
            start, end = anchor['start_char'], anchor['end_char']
            concordance.append([annotation['id'], index + 1, identity, name, label, annotation_name(result, annotation, 'functions'),
                                annotation['status'], segment['text'][max(0, start - 80):start] if current else '',
                                anchor['quote'], segment['text'][end:end + 80] if current else '',
                                segment['start'] if annotation.get('spans') else annotation['start'],
                                segment['end'] if annotation.get('spans') else annotation['end'], annotation.get('note', '')])
    paths = [folder / 'concordance.csv', folder / 'annotation-counts.csv', folder / 'speaker-summary.csv', folder / 'counting-rules.json']
    _write_csv(paths[0], ['id', 'span', 'speaker_id', 'speaker', 'label', 'function', 'status', 'left_context', 'quote', 'right_context', 'start', 'end', 'note'], concordance)
    _write_csv(paths[1], ['speaker_id', 'speaker', 'label', 'status', 'count'], [(*key, count) for key, count in sorted(counts.items())])
    _write_csv(paths[2], ['speaker_id', 'speaker_name', 'turns', 'words', 'sum_turn_seconds'],
               [[row['speaker'], row['name'], row['turns'], row['words'], row['seconds']] for row in speaker_statistics(result)])
    paths[3].write_text(json.dumps({'annotations': 'One record counts once per distinct source speaker. A multi-speaker annotation contributes to each involved speaker. Multiple labels count separately; statuses remain separate.',
                                  'words': 'Unicode word sequences, permitting internal apostrophes. Punctuation is excluded.',
                                  'duration': 'Sum of source turn durations. Overlapping turns count for each speaker, so the same time can be added twice.',
                                  'selection': result.get('metadata', {}).get('export_selection', {})}, indent=2), encoding='utf-8')
    return paths
