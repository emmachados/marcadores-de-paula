"""Exports from a common, editable segment representation. Times are seconds."""
from __future__ import annotations

import csv
import html
import json
import math
import re
from pathlib import Path
from xml.etree import ElementTree as ET

FORMATS = ("TXT", "DOCX", "CSV", "JSON", "SRT", "VTT", "TextGrid", "EAF")


def timestamp(seconds: float, decimal: str = ",") -> str:
    ms = max(0, round(float(seconds) * 1000))
    hours, ms = divmod(ms, 3600000)
    minutes, ms = divmod(ms, 60000)
    seconds, ms = divmod(ms, 1000)
    return f"{hours:02}:{minutes:02}:{seconds:02}{decimal}{ms:03}"


def speaker_name(speaker: str | None, names: dict) -> str:
    if not speaker:
        return "Unassigned"
    default = f"Speaker {int(speaker[8:]) + 1}" if speaker.startswith("SPEAKER_") and speaker[8:].isdigit() else speaker
    return names.get(speaker, default)


def clean_segments(result: dict) -> list[dict]:
    segments = []
    for s in result.get("segments", []):
        start, end = float(s["start"]), float(s["end"])
        if not math.isfinite(start + end) or start < 0 or end < start:
            raise ValueError("Transcript contains invalid timestamps.")
        segments.append({**s, "start": start, "end": max(end, start + .001)})
    return sorted(segments, key=lambda s: (s["start"], s["end"]))


def speaker_turns(result: dict) -> list[dict]:
    """Split aligned segments at word-level speaker changes, preserving every word."""
    turns = []
    for segment in result.get("segments", []):
        words = segment.get("words", [])
        if not words or not any(w.get("speaker") for w in words):
            turns.append(dict(segment))
            continue
        current = None
        source_text, cursor = segment.get("text", ""), 0
        first_turn = len(turns)
        for word in words:
            speaker = word.get("speaker")  # Uncertain words stay unassigned.
            start = word.get("start", current["end"] if current else segment["start"])
            end = word.get("end", start)
            if current is None or current.get("speaker") != speaker:
                current = {"start": start, "end": end, "speaker": speaker, "text": "", "words": []}
                turns.append(current)
            current["words"].append(dict(word))
            current["end"] = max(current["end"], end)
            word_text = word["word"].strip()
            position = source_text.find(word_text, cursor)
            if position >= 0:
                prefix = source_text[cursor:position]
                current["text"] += (prefix if current["text"] else prefix.lstrip()) + word_text
                cursor = position + len(word_text)
            else:
                # If normalisation changed the word strings, retain the original
                # segment instead of duplicating or silently changing its text.
                del turns[first_turn:]
                turns.append({**segment, "speaker_turn_split": "word_text_mismatch"})
                current = None
                break
        if current and cursor < len(source_text):
            current["text"] += source_text[cursor:]
    return turns


def _quote(text: str) -> str:
    return '"' + text.replace('"', '""').replace("\n", " ") + '"'


def _textgrid(segments, names, duration):
    tiers = {}
    for s in segments:
        tiers.setdefault(s.get("speaker"), []).append(s)
    lines = ['File type = "ooTextFile"', 'Object class = "TextGrid"', "", "xmin = 0",
             f"xmax = {duration:.6f}", "tiers? <exists>", f"size = {len(tiers)}", "item []:"]
    for i, (speaker, entries) in enumerate(tiers.items(), 1):
        intervals, cursor = [], 0.0
        for s in entries:
            if s["start"] < cursor - .000001:
                raise ValueError("Overlapping intervals within a speaker cannot be represented in one TextGrid tier. Use EAF or JSON.")
            start = max(cursor, s["start"])
            if start > cursor:
                intervals.append((cursor, start, ""))
            intervals.append((start, s["end"], s["text"]))
            cursor = s["end"]
        if cursor < duration:
            intervals.append((cursor, duration, ""))
        lines.extend([f"    item [{i}]:", '        class = "IntervalTier"',
                      f"        name = {_quote(speaker_name(speaker, names))}", "        xmin = 0",
                      f"        xmax = {duration:.6f}", f"        intervals: size = {len(intervals)}"])
        for j, (start, end, text) in enumerate(intervals, 1):
            lines.extend([f"        intervals [{j}]:", f"            xmin = {start:.6f}",
                          f"            xmax = {end:.6f}", f"            text = {_quote(text)}"])
    return "\n".join(lines) + "\n"


def _eaf(segments, names, metadata):
    root = ET.Element("ANNOTATION_DOCUMENT", AUTHOR="", DATE=metadata.get("created_at", ""),
                      FORMAT="3.0", VERSION="3.0",
                      attrib={"xmlns:xsi": "http://www.w3.org/2001/XMLSchema-instance",
                              "xsi:noNamespaceSchemaLocation": "http://www.mpi.nl/tools/elan/EAFv3.0.xsd"})
    header = ET.SubElement(root, "HEADER", MEDIA_FILE="", TIME_UNITS="milliseconds")
    media = metadata.get("review_media_path") or metadata.get("media_path")
    if media and Path(media).exists():
        import mimetypes
        ET.SubElement(header, "MEDIA_DESCRIPTOR", MEDIA_URL=Path(media).resolve().as_uri(),
                      MIME_TYPE=mimetypes.guess_type(media)[0] or "application/octet-stream")
    order = ET.SubElement(root, "TIME_ORDER")
    times = sorted({round(s[k] * 1000) for s in segments for k in ("start", "end")})
    time_ids = {t: f"ts{i}" for i, t in enumerate(times, 1)}
    for t, tid in time_ids.items():
        ET.SubElement(order, "TIME_SLOT", TIME_SLOT_ID=tid, TIME_VALUE=str(t))
    tiers = {}
    for i, s in enumerate(segments, 1):
        speaker = s.get("speaker") or "unassigned"
        if speaker not in tiers:
            tiers[speaker] = ET.SubElement(root, "TIER", TIER_ID=speaker,
                                          PARTICIPANT=speaker_name(s.get("speaker"), names),
                                          LINGUISTIC_TYPE_REF="utterance")
        ann = ET.SubElement(tiers[speaker], "ANNOTATION")
        aligned = ET.SubElement(ann, "ALIGNABLE_ANNOTATION", ANNOTATION_ID=f"a{i}",
                                TIME_SLOT_REF1=time_ids[round(s["start"] * 1000)],
                                TIME_SLOT_REF2=time_ids[round(s["end"] * 1000)])
        ET.SubElement(aligned, "ANNOTATION_VALUE").text = s["text"]
    ET.SubElement(root, "LINGUISTIC_TYPE", LINGUISTIC_TYPE_ID="utterance", TIME_ALIGNABLE="true", GRAPHIC_REFERENCES="false")
    ET.indent(root)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def export_result(result: dict, folder: Path, formats=FORMATS, stem="transcript", *, options=None) -> list[Path]:
    if options:
        from transcript_desk.export_options import selected_result
        result = selected_result(result, options)
    folder.mkdir(parents=True, exist_ok=True)
    segments = clean_segments(result)
    names, metadata = result.get("speaker_names", {}), result.get("metadata", {})
    duration = max(float(metadata.get("duration", 0)), max((s["end"] for s in segments), default=0))
    if options and options.get('annotation_tiers'):
        duration = max(duration, max((a['end'] for a in result.get('annotations', [])), default=0))
    written = []
    for fmt in formats:
        if fmt not in FORMATS:
            raise ValueError(f"Unsupported export: {fmt}")
        suffix = "TextGrid" if fmt == "TextGrid" else fmt.lower()
        path = folder / f"{stem}.{suffix}"
        if fmt == "JSON":
            path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        elif fmt == "TXT":
            if options:
                from transcript_desk.text_tools import reading_view
                text, _ = reading_view(result, timestamps=options.get('timestamps', True),
                                       paragraphs=options.get('paragraphs', True), grouped=options.get('grouped', False))
            else:
                text = "\n\n".join(f"[{timestamp(s['start'], '.')} – {timestamp(s['end'], '.')}] "
                                    f"{speaker_name(s.get('speaker'), names)}\n{s['text']}" for s in segments)
            path.write_text(text + "\n", encoding="utf-8")
        elif fmt == "CSV":
            with path.open("w", encoding="utf-8-sig", newline="") as file:
                writer = csv.writer(file)
                writer.writerow(["start_seconds", "end_seconds", "speaker_id", "speaker_name", "text"])
                for s in segments:
                    # Avoid spreadsheet formula execution in cells containing source text.
                    safe = lambda v: "'" + v if str(v).lstrip().startswith(("=", "+", "-", "@")) else v
                    writer.writerow([s["start"], s["end"], s.get("speaker", ""), safe(speaker_name(s.get("speaker"), names)), safe(s["text"])])
        elif fmt in ("SRT", "VTT"):
            blocks = []
            for i, s in enumerate(segments, 1):
                text = f"{speaker_name(s.get('speaker'), names)}: {s['text']}"
                # Escape caption markup, preserve Unicode, prevent blank-line cue termination.
                text = html.escape(re.sub(r"\s+", " ", text).strip(), quote=False)
                decimal = "." if fmt == "VTT" else ","
                blocks.append(f"{i}\n{timestamp(s['start'], decimal)} --> {timestamp(s['end'], decimal)}\n{text}")
            path.write_text(("WEBVTT\n\n" if fmt == "VTT" else "") + "\n\n".join(blocks) + "\n", encoding="utf-8")
        elif fmt == "TextGrid":
            if options and options.get('annotation_tiers'):
                from transcript_desk.export_options import annotation_tier_segments
                tiers, tier_names = annotation_tier_segments(result)
                path.write_text(_textgrid(segments + tiers, {**names, **tier_names}, duration), encoding="utf-8")
            else:
                path.write_text(_textgrid(segments, names, duration), encoding="utf-8")
        elif fmt == "EAF":
            if options and options.get('annotation_tiers'):
                from transcript_desk.export_options import annotation_tier_segments
                tiers, tier_names = annotation_tier_segments(result)
                path.write_bytes(_eaf(segments + tiers, {**names, **tier_names}, metadata))
            else:
                path.write_bytes(_eaf(segments, names, metadata))
        elif fmt == "DOCX":
            from docx import Document
            doc = Document()
            doc.add_heading(metadata.get("title", "Transcript"), 0)
            doc.add_paragraph(f"Source: {metadata.get('source', '')}")
            if options:
                from transcript_desk.export_options import add_document_content
                add_document_content(doc, result, options)
            else:
                for s in segments:
                    doc.add_paragraph(f"{timestamp(s['start'], '.')}  {speaker_name(s.get('speaker'), names)}", style="Heading 2")
                    doc.add_paragraph(s["text"])
            doc.save(path)
        written.append(path)
    if result.get("annotations"):
        from transcript_desk.annotations import export_annotations
        written.extend(export_annotations(result, folder))
    if options and options.get('analysis_tables'):
        from transcript_desk.export_options import export_analysis
        written.extend(export_analysis(result, folder))
    return written
