"""End-to-end checks in the shipped interface with synthetic review content."""
import copy
import json
import math
import struct
import traceback
import wave
from pathlib import Path

from PySide6.QtCore import QTimer, Qt
from PySide6.QtTest import QTest


def run_workspace_check(application, window, report_path):
    report = Path(report_path).resolve()
    report.parent.mkdir(parents=True, exist_ok=True)
    from transcript_desk.runtime import data_root
    from transcript_desk.exports import speaker_name
    from transcript_desk.review_storage import atomic_write, load_review
    from transcript_desk.annotations import create_annotation, create_range_annotation, create_time_annotation
    from transcript_desk.codebook import new_entry, ensure_codebook
    from transcript_desk.project_archive import save_snapshot, write_archive, read_archive, list_snapshots
    from transcript_desk.workspace_panel import SearchDialog, commit_review
    from transcript_desk.export_dialog import ExportDialog
    from transcript_desk.codebook_dialog import CodebookDialog
    folder = data_root() / 'Synthetic workspace verification'
    folder.mkdir(parents=True, exist_ok=True)
    audio = folder / 'audio.wav'
    with wave.open(str(audio), 'wb') as stream:
        stream.setparams((1, 2, 16000, 0, 'NONE', 'not compressed'))
        stream.writeframes(b''.join(struct.pack('<h', int(3500 * math.sin(2 * math.pi * 220 * i / 16000))) for i in range(16000 * 6)))
    result = {'metadata': {'title': 'Synthetic annotation and speaker review', 'source': 'Synthetic verification fixture',
                           'duration': 6, 'media_path': str(audio)}, 'speaker_names': {'A': 'Paula', 'B': 'Participant'},
              'segments': [{'start': 0, 'end': 1.5, 'speaker': 'A', 'text': 'Well, shall we begin with the first question?'},
                           {'start': 1.6, 'end': 3, 'speaker': 'B', 'text': 'Yes, I I think that example is clear.'},
                           {'start': 3.1, 'end': 4.5, 'speaker': 'A', 'text': 'Well, could you describe the next example?'},
                           {'start': 4.6, 'end': 6, 'speaker': 'B', 'text': 'There is a brief pause, then a correction.'}]}
    result['annotations'] = [create_annotation(result, 0, 0, 4, 'Discourse marker'),
                             create_annotation(result, 1, 7, 10, 'Self-repetition', status='suggested'),
                             create_range_annotation(result, 2, 3, 'Topic boundary'),
                             create_time_annotation(result, 3, 4.7, 4.85, 'Breath')]
    book = ensure_codebook(result)
    entry = new_entry('Project observation')
    entry['definition'] = 'A synthetic category used to verify project codebook persistence.'
    book['labels'].append(entry)
    atomic_write(folder / 'result.json', result)
    status = {'synthetic_fixture': True, 'checks': []}

    def checked(name):
        status['checks'].append(name)

    def fail(exc):
        status.update(ok=False, error=str(exc), traceback=traceback.format_exc())
        report.write_text(json.dumps(status, indent=2), encoding='utf-8')
        window.workspace.audio.shutdown()
        application.exit(1)

    def stage_one():
        try:
            assert window.open_project_folder(folder)
            window.player.audio.setMuted(True)
            # The final name matches speaker A, returning the turn to that identity.
            for iteration, name in enumerate([f'Paula {i}' for i in range(12)] + ['Paula']):
                window.transcript.setCurrentCell(0, 2)
                window.transcript.editItem(window.transcript.item(0, 2))
                application.processEvents()
                editor = window.delegate.editor
                editor.setText(name)
                window.highlight_position(700)
                QTest.keyClick(editor, Qt.Key_Return)
                application.processEvents()
                window.finish_edit()
                assert speaker_name(window.result['segments'][0]['speaker'], window.result['speaker_names']) == name
            assert window.result['segments'][0]['speaker'] == 'A' and window.result['speaker_names']['A'] == 'Paula'
            checked('13 live speaker commits with playback highlighting')
            revised = copy.deepcopy(window.result)
            revised['speaker_names']['A'] = 'Paula'
            commit_review(window, revised)
            search = SearchDialog(window)
            search.query.setText('Well')
            search.replacement.setText('Now')
            assert len(search.matches) == 2
            search.replace_all()
            assert window.result['segments'][2]['text'].startswith('Now')
            window.move_review_history(-1)
            assert window.result['segments'][2]['text'].startswith('Well')
            search.deleteLater()
            checked('previewed replacement and one-step undo')
            assert window.apply_turn_edit('split', 2, 6, 3.5)
            assert window.apply_turn_edit('merge', 2)
            checked('split/merge with multi-turn annotation')
            window.tabs.setCurrentWidget(window.workspace.reading)
            reading = window.workspace.reading
            reading.speaker.setCurrentIndex(reading.speaker.findData('A'))
            assert 'first question' in reading.text.toPlainText() and 'next example' in reading.text.toPlainText()
            assert '00:00' not in reading.text.toPlainText() and 'I I think' not in reading.text.toPlainText()
            application.processEvents()
            window.grab().save(str(report.with_name(report.stem + '-speaker-text.png')))
            checked('interleaved speaker grouping without timestamps')
            window.tabs.setCurrentWidget(window.annotations)
            window.annotations.bind(window.result)
            application.processEvents()
            window.grab().save(str(report.with_name(report.stem + '-annotations.png')))
            assert window.save_edits()
            save_snapshot(folder, window.result, 'Verified annotation review')
            archive = write_archive(folder, window.result, folder.parent / 'portable.paula.zip')
            imported = read_archive(archive, folder.parent / 'Imported portable project')
            assert list_snapshots(imported)
            from transcript_desk.projects import media_in_project
            assert media_in_project(load_review(imported)[0], imported).is_file()
            checked('portable project including media and named snapshot')
            export = ExportDialog(window)
            export.preset.setCurrentIndex(3)
            export.show()
            application.processEvents()
            export.grab().save(str(report.with_name(report.stem + '-export.png')))
            options = export.options()
            destination = window.workspace.export_to(folder / 'exports', ['TXT', 'DOCX', 'CSV', 'JSON', 'SRT', 'VTT', 'TextGrid', 'EAF'], options)
            assert destination and (destination / 'annotation-counts.csv').exists()
            export.close()
            export.deleteLater()
            checked('eight transcript exports with filtered annotation layers and counts')
            for key, action in window.layout_actions.items():
                action.setChecked(key == 'grouped')
            window.result_job['config']['text_layout'] = window.options()['text_layout']
            export = ExportDialog(window)
            assert (export.times.isChecked(), export.grouped.isChecked(), export.paragraphs.isChecked()) == (False, True, False)
            export.grouped.setChecked(False)
            destination = window.workspace.export_to(folder / 'exports-layout', ['TXT', 'DOCX'], export.options())
            export.deleteLater()
            from docx import Document
            headings = [p.text for p in Document(destination / 'transcript.docx').paragraphs if p.style.name.startswith('Heading')]
            rows = sorted(window.result['segments'], key=lambda s: (s['start'], s['end']))
            assert len(headings) == sum(i == 0 or s.get('speaker') != rows[i - 1].get('speaker') for i, s in enumerate(rows))
            assert '00:00' not in (destination / 'transcript.txt').read_text(encoding='utf-8')
            checked('recording text layout reaches the export switches; DOCX headings at speaker changes')
            window.tabs.setCurrentIndex(0)
            window.workspace.audio.expand.setChecked(True)
            window.workspace.audio.load_waveform()
            QTimer.singleShot(100, wait_waveform)
        except Exception as exc:
            fail(exc)

    attempts = [0]

    def wait_waveform():
        try:
            attempts[0] += 1
            if window.workspace.audio.process is not None:
                if attempts[0] > 150:
                    raise RuntimeError('Waveform decode timed out')
                QTimer.singleShot(100, wait_waveform)
                return
            assert len(window.workspace.audio.wave.peaks) >= 590
            checked('bundled FFmpeg waveform decoding')
            window.workspace.audio.play_range(.2, .8)
            QTimer.singleShot(1200, finish)
        except Exception as exc:
            fail(exc)

    def finish():
        try:
            assert window.workspace.audio.active_interval is None
            checked('actual audio interval playback stops at boundary')
            window.grab().save(str(report.with_name(report.stem + '-waveform.png')))
            assert window.save_edits()
            window.load_result(window.result_job)
            assert len(window.result['annotations']) == 4
            assert any(e['name'] == 'Project observation' for e in window.result['annotation_codebook']['labels'])
            checked('saved annotations and custom codebook reopen')
            status.update(ok=True, version=__import__('transcript_desk').__version__)
            report.write_text(json.dumps(status, indent=2), encoding='utf-8')
            window.close()
            application.exit(0)
        except Exception as exc:
            fail(exc)

    QTimer.singleShot(300, stage_one)
